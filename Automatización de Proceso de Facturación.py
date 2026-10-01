import openpyxl
import os
import smtplib
from email.message import EmailMessage
import mimetypes


# ========================================
ruta_excel = r"Y:\Tableros 2026\Staffing\2026 Tablero Control Staffing - GENERAL.xlsx"

# credenciales de Gmail
correo_emisor = "sbarbas@sms-sudamerica.com"
password_emisor = "hqrtxtudgoeabedl"

# destinatarios
correo_receptor = "adm@sms-sudamerica.com"
correos_cc = "staffing_arg@sms-sudamerica.com", "fmenchacabaso@sms-sudamerica.com", "direccion@sms-sudamerica.com", "jortiz@sms-sudamerica.com"

# =========================================


print("Abriendo Excel (esto puede demorar unos segundos)...")
# abre el Excel dos veces (una para lectura y otra para escritura)
wb_valores = openpyxl.load_workbook(ruta_excel, data_only=True)
ws_valores = wb_valores["Fact_TOTAL"]

wb_guardar = openpyxl.load_workbook(ruta_excel)
ws_guardar = wb_guardar["Fact_TOTAL"]

mails_a_enviar = {}

nombres_meses = {
    "01": "ENERO", "02": "FEBRERO", "03": "MARZO", "04": "ABRIL",
    "05": "MAYO", "06": "JUNIO", "07": "JULIO", "08": "AGOSTO",
    "09": "SEPTIEMBRE", "10": "OCTUBRE", "11": "NOVIEMBRE", "12": "DICIEMBRE"
}


for row in range(2, ws_valores.max_row + 1):
    estado = ws_valores.cell(row=row, column=20).value  # Columna T (ESTADO)

    if estado == "Apto Facturar":
        cliente = ws_valores.cell(row=row, column=3).value   # Columna C (Para agrupar por cliente)
        adjunto = ws_valores.cell(row=row, column=17).value  # Columna Q (Documento a adjuntar)

        def obtener_valor(col):
            val = ws_valores.cell(row=row, column=col).value
            return str(val).strip() if val is not None and str(val).lower() != "nan" and str(val).strip() != "" else None

        def formatear_moneda(valor_raw, tipo_moneda):
            if not valor_raw:
                return None
            try:
                monto_num = float(valor_raw)
                monto_formateado = f"{monto_num:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
            except ValueError:
                monto_formateado = valor_raw
            
            if tipo_moneda == "Dólares":
                return f"USD {monto_formateado}"
            elif tipo_moneda == "Pesos":
                return f"$ {monto_formateado}"
            else:
                return monto_formateado

        razon_social = obtener_valor(6) or "S/D"
        mes = obtener_valor(7) or "S/D"

        if cliente not in mails_a_enviar:
            mails_a_enviar[cliente] = {
                'filas_excel': [],
                'adjuntos_unicos': set(),
                'cuerpo_texto': "",
                'razon_social_asunto': razon_social,
                'mes_asunto': mes
            }

        # guarda el número de la fila para después escribir el "Si"
        mails_a_enviar[cliente]['filas_excel'].append(row)

        # evita duplicados si es el mismo archivo
        if adjunto and str(adjunto).strip() != "" and str(adjunto).lower() != "nan":
            if os.path.exists(str(adjunto)):
                mails_a_enviar[cliente]['adjuntos_unicos'].add(str(adjunto))
            else:
                print(f"⚠️ Omitido (No encontrado en tu PC o VPN): {adjunto}")


        # -- Datos obligatorios --
        proyecto = obtener_valor(2) or "S/D"
        servicio_recurso = obtener_valor(4) or "S/D"
        
        moneda = obtener_valor(8)    
        monto_raw = obtener_valor(9) 
        monto_final = formatear_moneda(monto_raw, moneda) or "S/D"


        bloque_fila = f"  PROYECTO: {proyecto}\n  SERVICIO-RECURSO: {servicio_recurso}\n  AÑO/MES: {mes}\n  MONTO: {monto_final}\n"


        # -- Datos opcionales --
        # se imprimen en el mail sólo si la celda tiene algo escrito)

        adicionales_final = formatear_moneda(obtener_valor(14), moneda)

        opcionales = {
            "OC": obtener_valor(10),                 # Col J
            "RECEPCIÓN": obtener_valor(11),          # Col K
            "REQUIRENTE": obtener_valor(12),         # Col L
            "TARIFA-HORA": obtener_valor(13),        # Col M
            "ADICIONALES-BONOS": adicionales_final,  # Col N
            "CENTRO DE COSTO": obtener_valor(15),    # Col O
            "CUIT": obtener_valor(16),               # Col P
            "OBSERVACIONES": obtener_valor(19)       # Col S
        }

        for etiqueta, valor in opcionales.items():
            if valor:
                bloque_fila += f"  {etiqueta}: {valor}\n"

        bloque_fila += "-" * 50 + "\n"
        mails_a_enviar[cliente]['cuerpo_texto'] += bloque_fila


# envía a través de Gmail y actualiza el archivo
if mails_a_enviar:
    try:
        print("\nConectando con Gmail...")
        server = smtplib.SMTP('smtp.gmail.com', 587)
        server.starttls()
        server.login(correo_emisor, password_emisor)

        for cliente, datos in mails_a_enviar.items():
            msg = EmailMessage()
            msg['From'] = correo_emisor
            msg['To'] = correo_receptor

            if correos_cc.strip():
                msg['Cc'] = correos_cc


            # -- Asunto del mail --

            mes_raw = datos['mes_asunto']
            
            if "-" in mes_raw and len(mes_raw) == 7:
                anio, mes_num = mes_raw.split("-")
                mes_formateado = f"{nombres_meses.get(mes_num, mes_num)} {anio}"
            else:
                mes_formateado = mes_raw # si el formato es distinto, lo deja como está
                
            msg['Subject'] = f"Facturación Servicios {mes_formateado} - {datos['razon_social_asunto']}"

            
            # -- Mensaje del mail --

            cuerpo_completo = f"Buenos días Equipo,\nLes pasamos lo que se puede facturar de {cliente} por el mes de {mes}:\n\n{datos['cuerpo_texto']}\n\n¡Saludos y Gracias!"
            msg.set_content(cuerpo_completo)

            # -- Carga archivos adjuntos --
            for ruta_adjunto in datos['adjuntos_unicos']:
                tipo_mime, _ = mimetypes.guess_type(ruta_adjunto)
                if tipo_mime is None:
                    tipo_mime = 'application/octet-stream'
                tipo_principal, sub_tipo = tipo_mime.split('/', 1)

                with open(ruta_adjunto, 'rb') as f:
                    msg.add_attachment(f.read(), maintype=tipo_principal, subtype=sub_tipo, filename=os.path.basename(ruta_adjunto))

            # envía el correo
            server.send_message(msg)
            print(f"✅ Mail enviado exitosamente para {cliente}")

            # actualiza el Excel con el "Si" en la columna U (solo para este cliente)
            for fila in datos['filas_excel']:
                ws_guardar.cell(row=fila, column=21).value = "Si"

        server.quit()

        print("Guardando Excel...")
        wb_guardar.save(ruta_excel)
        print("\n✅ Proceso 100% finalizado. Excel guardado y actualizado.")

    except Exception as e:
        print(f"\n❌ Ocurrió un error con Gmail o guardando el archivo: {e}")
else:
    print("\nNo se encontraron filas con el estado 'Apto Facturar'. No se envió nada.")

input("\nPresiona Enter para cerrar la ventana...")